"""P1: exports (CSV/XLSX/PDF), public results + medal table, QR certificates, audit log."""
from __future__ import annotations
import csv, io, json, secrets
from fastapi import APIRouter, Depends, HTTPException, Response
from sqlalchemy.orm import Session
from app.core.db import get_db
from app.core.deps import require_perm, require_tournament_owner
from app.core.paging import page_args, paginate, envelope
from app.models.user import User
from app.models.club_athlete import Athlete, Club
from app.models.tournament import Tournament, TournamentCategory
from app.models.competition import Registration, Bracket, BracketMatch
from app.models.misc import AuditLog, Document

router = APIRouter(tags=["exports"])

# ---------- results / medals (§47) ----------

def _category_standings(db: Session, tid: int) -> list[dict]:
    brackets = db.query(Bracket).filter_by(tournament_id=tid).all()
    if not brackets:
        return []
    bids = [b.id for b in brackets]
    cmap = {c.id: c for c in db.query(TournamentCategory).filter(
        TournamentCategory.id.in_([b.category_id for b in brackets])).all()}
    finals = db.query(BracketMatch).filter(
        BracketMatch.bracket_id.in_(bids),
        BracketMatch.status == "finished",
        BracketMatch.next_match_id.is_(None)).all()
    by_bracket: dict[int, list] = {}
    for m in finals:
        by_bracket.setdefault(m.bracket_id, []).append(m)
    aids = {m.winner_id for m in finals} | {m.athlete_a_id for m in finals} | {m.athlete_b_id for m in finals}
    aids.discard(None)
    amap = {a.id: a for a in db.query(Athlete).filter(Athlete.id.in_(aids)).all()} if aids else {}
    def name(aid):
        a = amap.get(aid) if aid else None
        return a.full_name if a else "—"
    out = []
    for b in brackets:
        fin = by_bracket.get(b.id, [])
        cat = cmap.get(b.category_id)
        winner = fin[0].winner_id if fin else None
        finalists = {m.athlete_a_id for m in fin} | {m.athlete_b_id for m in fin}
        out.append({"category_id": b.category_id, "category": cat.name if cat else "?",
                    "champion_id": winner, "champion": name(winner),
                    "finalists": [name(f) for f in finalists if f and f != winner]})
    return out

@router.get("/api/tournaments/{tid}/results")
def results(tid: int, db: Session = Depends(get_db)):
    t = db.get(Tournament, tid)
    if not t:
        raise HTTPException(404, "Not found")
    standings = _category_standings(db, tid)
    medals: dict[int, dict] = {}
    for s in standings:
        if s["champion_id"]:
            m = medals.setdefault(s["champion_id"], {"athlete_id": s["champion_id"], "gold": 0, "titles": []})
            m["gold"] += 1
            m["titles"].append(s["category"])
    table = sorted(medals.values(), key=lambda m: -m["gold"])
    aids = [m["athlete_id"] for m in table]
    amap = {a.id: a for a in db.query(Athlete).filter(Athlete.id.in_(aids)).all()} if aids else {}
    cids = [a.club_id for a in amap.values() if a.club_id]
    clubmap = {c.id: c for c in db.query(Club).filter(Club.id.in_(cids)).all()} if cids else {}
    for m in table:
        a = amap.get(m["athlete_id"])
        m["athlete"] = a.full_name if a else "?"
        club = clubmap.get(a.club_id) if a and a.club_id else None
        m["club"] = club.name if club else "—"
        # B2: club profile deep-link (additive key, public name already shown).
        m["club_id"] = club.id if club else None
    return {"tournament": t.name, "standings": standings, "medal_table": table}

# ---------- exports (§21) ----------

def _csv_safe(v) -> str:
    """P1: neutralize CSV formula injection — names/clubs are user input and
    Excel executes cells starting with =+-@. Prefixing with ' keeps data
    intact while blocking execution."""
    s = "" if v is None else str(v)
    return "'" + s if s[:1] in ("=", "+", "-", "@") else s

def _participants(db: Session, tid: int) -> list[dict]:
    regs = db.query(Registration).filter_by(tournament_id=tid).all()
    if not regs:
        return []
    amap = {a.id: a for a in db.query(Athlete).filter(
        Athlete.id.in_([r.athlete_id for r in regs])).all()}
    cmap = {c.id: c for c in db.query(TournamentCategory).filter(
        TournamentCategory.id.in_([r.category_id for r in regs])).all()}
    cids = [a.club_id for a in amap.values() if a.club_id]
    clubmap = {c.id: c for c in db.query(Club).filter(Club.id.in_(cids)).all()} if cids else {}
    rows = []
    for r in regs:
        a = amap.get(r.athlete_id)
        c = cmap.get(r.category_id)
        club = clubmap.get(a.club_id) if a and a.club_id else None
        rows.append({"athlete": a.full_name if a else "?", "club": club.name if club else "—",
                     "country": a.country if a else "", "category": c.name if c else "?",
                     "weight": r.weigh_in_kg or (a.weight_kg if a else ""),
                     "weigh_in": r.weigh_in_status, "checked_in": "yes" if r.checked_in else "no"})
    return rows

@router.get("/api/tournaments/{tid}/export/participants.csv")
def exp_part_csv(tid: int, db: Session = Depends(get_db), user: User = Depends(require_perm("tournaments.manage"))):
    require_tournament_owner(tid, db, user)
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=["athlete", "club", "country", "category", "weight", "weigh_in", "checked_in"])
    w.writeheader()
    w.writerows({k: (_csv_safe(v) if isinstance(v, str) else v) for k, v in r.items()} for r in _participants(db, tid))
    return Response(buf.getvalue(), media_type="text/csv",
                    headers={"Content-Disposition": f"attachment; filename=t{tid}-participants.csv"})

@router.get("/api/tournaments/{tid}/export/participants.xlsx")
def exp_part_xlsx(tid: int, db: Session = Depends(get_db), user: User = Depends(require_perm("tournaments.manage"))):
    require_tournament_owner(tid, db, user)
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Participants"
    ws.append(["Athlete", "Club", "Country", "Category", "Weight", "Weigh-in", "Checked-in"])
    for r in _participants(db, tid):
        ws.append([r["athlete"], r["club"], r["country"], r["category"], r["weight"], r["weigh_in"], r["checked_in"]])
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename=t{tid}-participants.xlsx"})

@router.get("/api/tournaments/{tid}/export/schedule.csv")
def exp_sched_csv(tid: int, db: Session = Depends(get_db), user: User = Depends(require_perm("tournaments.manage"))):
    require_tournament_owner(tid, db, user)
    bids = [b.id for b in db.query(Bracket).filter_by(tournament_id=tid).all()]
    ms = db.query(BracketMatch).filter(BracketMatch.bracket_id.in_(bids)).order_by(BracketMatch.scheduled_at).all() if bids else []
    aids = ({m.athlete_a_id for m in ms} | {m.athlete_b_id for m in ms}) - {None}
    amap = {a.id: a for a in db.query(Athlete).filter(Athlete.id.in_(aids)).all()} if aids else {}
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["match_id", "round", "athlete_a", "athlete_b", "tatami", "scheduled_at", "status"])
    for m in ms:
        def nm(aid):
            a = amap.get(aid) if aid else None
            return _csv_safe(a.full_name) if a else "bye"
        w.writerow([m.id, m.round_no, nm(m.athlete_a_id), nm(m.athlete_b_id), m.tatami_id, m.scheduled_at, m.status])
    return Response(buf.getvalue(), media_type="text/csv",
                    headers={"Content-Disposition": f"attachment; filename=t{tid}-schedule.csv"})

def _pdf(title: str, lines: list[str]) -> bytes:
    from reportlab.pdfgen import canvas
    from reportlab.lib.pagesizes import A4
    from app.services.pdf_fonts import ensure_fonts
    regular, bold = ensure_fonts()
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    c.setFont(bold, 18)
    c.drawString(60, 780, title)
    c.setFont(regular, 11)
    y = 750
    for ln in lines:
        c.drawString(60, y, ln[:110])
        y -= 18
        if y < 60:
            c.showPage()
            y = 780
    c.showPage()
    c.save()
    return buf.getvalue()

@router.get("/api/tournaments/{tid}/export/protocol.pdf")
def exp_protocol(tid: int, db: Session = Depends(get_db), user: User = Depends(require_perm("tournaments.manage"))):
    require_tournament_owner(tid, db, user)
    t = db.get(Tournament, tid)
    if not t:
        raise HTTPException(404, "Not found")
    lines = [f"{s['category']}: champion {s['champion']}" for s in _category_standings(db, tid)]
    pdf = _pdf(f"Protocol - {t.name}", lines or ["No results yet"])
    return Response(pdf, media_type="application/pdf",
                    headers={"Content-Disposition": f"attachment; filename=t{tid}-protocol.pdf"})

# ---------- Wave 1: print sheets (weigh-in / start protocol / tatami schedule) ----------

def _sheet_pdf(title: str, headers: list[str], rows: list[list]) -> bytes:
    """One A4 table helper on bundled Cyrillic fonts."""
    from reportlab.pdfgen import canvas
    from reportlab.lib.pagesizes import A4
    from app.services.pdf_fonts import ensure_fonts
    regular, bold = ensure_fonts()
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    widths = [38, 200, 130, 70, 60, 70][:len(headers)]
    xs = [40]
    for w_ in widths[:-1]:
        xs.append(xs[-1] + w_)

    def header(y0):
        c.setFont(bold, 14)
        c.drawString(40, y0, title[:90])
        c.setFont(bold, 10)
        for x, h in zip(xs, headers):
            c.drawString(x, y0 - 22, h[:28])
        return y0 - 40

    y = header(800)
    c.setFont(regular, 10)
    for row in rows:
        if y < 60:
            c.showPage()
            y = header(800)
            c.setFont(regular, 10)
        for x, cell in zip(xs, row):
            c.drawString(x, y, str(cell if cell is not None else "")[:30])
        y -= 18
    c.showPage()
    c.save()
    return buf.getvalue()


@router.get("/api/tournaments/{tid}/export/weighin.pdf")
def exp_weighin_pdf(tid: int, db: Session = Depends(get_db), user: User = Depends(require_perm("tournaments.manage"))):
    """Wave 1: printable weigh-in sheet (athlete, category, weight, status)."""
    require_tournament_owner(tid, db, user)
    rows = []
    for r in _participants(db, tid):
        rows.append([r["athlete"], r["category"], r["weight"], r["weigh_in"], r["checked_in"]])
    pdf = _sheet_pdf("Weigh-in sheet", ["Athlete", "Category", "Weight", "Status", "In"], rows)
    return Response(pdf, media_type="application/pdf",
                    headers={"Content-Disposition": f"attachment; filename=t{tid}-weighin.pdf"})


@router.get("/api/tournaments/{tid}/export/start-protocol.pdf")
def exp_start_protocol_pdf(tid: int, db: Session = Depends(get_db), user: User = Depends(require_perm("tournaments.manage"))):
    """Wave 1: start protocol grouped by category in seed order."""
    require_tournament_owner(tid, db, user)
    regs = db.query(Registration).filter_by(tournament_id=tid).order_by(Registration.seed.desc()).all()
    amap = {a.id: a for a in db.query(Athlete).filter(
        Athlete.id.in_([r.athlete_id for r in regs])).all()} if regs else {}
    cmap = {c.id: c for c in db.query(TournamentCategory).filter_by(tournament_id=tid).all()}
    by_cat: dict[int, list] = {}
    for r in regs:
        by_cat.setdefault(r.category_id, []).append(r)
    rows = []
    for cid, rs in by_cat.items():
        c = cmap.get(cid)
        rows.append([f"== {c.name if c else '?'} ==", "", "", "", ""])
        for i, r in enumerate(rs, 1):
            a = amap.get(r.athlete_id)
            rows.append([i, a.full_name if a else "?", a.country if a else "",
                         a.weight_kg if a else "", ""])
    pdf = _sheet_pdf("Start protocol", ["#", "Athlete", "Country", "Weight", "Sign"], rows)
    return Response(pdf, media_type="application/pdf",
                    headers={"Content-Disposition": f"attachment; filename=t{tid}-start-protocol.pdf"})


@router.get("/api/tournaments/{tid}/export/schedule.pdf")
def exp_schedule_pdf(tid: int, db: Session = Depends(get_db), user: User = Depends(require_perm("tournaments.manage"))):
    """Wave 1: printable tatami schedule (match order, names, times)."""
    require_tournament_owner(tid, db, user)
    bids = [b.id for b in db.query(Bracket).filter_by(tournament_id=tid).all()]
    ms = db.query(BracketMatch).filter(BracketMatch.bracket_id.in_(bids)).order_by(
        BracketMatch.tatami_id, BracketMatch.scheduled_at).all() if bids else []
    aids = ({m.athlete_a_id for m in ms} | {m.athlete_b_id for m in ms}) - {None}
    amap = {a.id: a for a in db.query(Athlete).filter(Athlete.id.in_(aids)).all()} if aids else {}
    rows = []
    for m in ms:
        def nm(aid):
            a = amap.get(aid) if aid else None
            return a.full_name if a else "bye"
        at = str(m.scheduled_at)[11:16] if m.scheduled_at else ""
        rows.append([m.id, f"T{m.tatami_id or '—'}", at, f"{nm(m.athlete_a_id)} — {nm(m.athlete_b_id)}", m.status])
    if not rows:
        rows = [["—", "", "", "No fights scheduled", ""]]
    pdf = _sheet_pdf("Tatami schedule", ["Match", "Tatami", "Time", "Fight", "Status"], rows)
    return Response(pdf, media_type="application/pdf",
                    headers={"Content-Disposition": f"attachment; filename=t{tid}-schedule.pdf"})

# ---------- documents / certificates (§19) ----------

def _podium(db: Session, tid: int) -> list[dict]:
    """Wave 1: places 1-3 per finished category from single-elimination data.

    Honest mapping (no bronze fight exists): gold = final winner,
    silver = final loser, bronze = semifinal losers. Unfinished categories
    yield no places. Returns [{category_id, category, gold_id, silver_id,
    bronze_ids}].
    """
    brackets = db.query(Bracket).filter_by(tournament_id=tid).all()
    if not brackets:
        return []
    bids = [b.id for b in brackets]
    cmap = {c.id: c for c in db.query(TournamentCategory).filter(
        TournamentCategory.id.in_([b.category_id for b in brackets])).all()}
    ms = db.query(BracketMatch).filter(BracketMatch.bracket_id.in_(bids)).all()
    by_bracket: dict[int, list] = {}
    for m in ms:
        by_bracket.setdefault(m.bracket_id, []).append(m)
    return _assemble_podium(brackets, cmap, by_bracket)


def _assemble_podium(brackets, cmap: dict, by_bracket: dict) -> list[dict]:
    """Pure podium math shared by _podium (single) and _podiums_bulk.
    Same mapping: gold = final winner, silver = final loser,
    bronze = semifinal losers (no bronze fight exists)."""
    out = []
    for b in brackets:
        matches = by_bracket.get(b.id, [])
        finals = [m for m in matches
                  if m.next_match_id is None and m.status == "finished" and m.winner_id]
        if not finals:
            continue
        fin = finals[0]
        loser = fin.athlete_a_id if fin.winner_id == fin.athlete_b_id else fin.athlete_b_id
        semis = [m for m in matches
                 if m.next_match_id == fin.id and m.status == "finished" and m.winner_id]
        bronze = []
        for s in semis:
            l = s.athlete_a_id if s.winner_id == s.athlete_b_id else s.athlete_b_id
            if l and l != s.winner_id:
                bronze.append(l)
        out.append({"category_id": b.category_id,
                    "category": cmap[b.category_id].name if b.category_id in cmap else "?",
                    "gold_id": fin.winner_id,
                    "silver_id": loser,
                    "bronze_ids": sorted(set(bronze) - {fin.winner_id, loser})})
    return out


def _podiums_bulk(db: Session, tids: list[int]) -> dict[int, list[dict]]:
    """Batched podiums for many tournaments (B2 club profile): 3 queries
    total regardless of len(tids) — no N+1. Same math as _podium."""
    out: dict[int, list[dict]] = {tid: [] for tid in tids}
    if not tids:
        return out
    brackets = db.query(Bracket).filter(Bracket.tournament_id.in_(tids)).all()
    if not brackets:
        return out
    cmap = {c.id: c for c in db.query(TournamentCategory).filter(
        TournamentCategory.id.in_([b.category_id for b in brackets])).all()}
    ms = db.query(BracketMatch).filter(
        BracketMatch.bracket_id.in_([b.id for b in brackets])).all()
    by_bracket: dict[int, list] = {}
    for m in ms:
        by_bracket.setdefault(m.bracket_id, []).append(m)
    by_tid: dict[int, list] = {}
    for b in brackets:
        by_tid.setdefault(b.tournament_id, []).append(b)
    for tid, bs in by_tid.items():
        out[tid] = _assemble_podium(bs, cmap, by_bracket)
    return out


def _issue_one(db: Session, athlete_id: int, tournament_id: int, kind: str,
               place: str, category: str) -> Document | None:
    """Create a diploma unless an identical one already exists (idempotent bulk)."""
    exists = db.query(Document).filter_by(athlete_id=athlete_id, tournament_id=tournament_id,
                                          kind=kind).first()
    if exists and kind == "diploma":
        # same athlete+kind: keep the best (lowest) place
        try:
            old = json.loads(exists.payload or "{}")
        except ValueError:
            old = {}
        try:
            if old.get("place") and int(old["place"]) <= int(place):
                return None
        except (ValueError, TypeError):
            pass
        exists.payload = json.dumps({"place": place[:32], "category": category[:128]})
        return exists
    if exists:
        return None
    d = Document(code=secrets.token_hex(16).upper(), kind=kind, athlete_id=athlete_id,
                 tournament_id=tournament_id,
                 payload=json.dumps({"place": place[:32], "category": category[:128]}))
    db.add(d)
    return d


@router.get("/api/tournaments/{tid}/podium")
def podium(tid: int, db: Session = Depends(get_db)):
    """Wave 1: public places 1-3 per finished category (names resolved)."""
    t = db.get(Tournament, tid)
    if not t:
        raise HTTPException(404, "Not found")
    pod = _podium(db, tid)
    aids = set()
    for p in pod:
        aids.add(p["gold_id"])
        aids.add(p["silver_id"])
        aids.update(p["bronze_ids"])
    aids.discard(None)
    amap = {a.id: a for a in db.query(Athlete).filter(Athlete.id.in_(list(aids))).all()} if aids else {}
    cids = [a.club_id for a in amap.values() if a.club_id]
    clubmap = {c.id: c for c in db.query(Club).filter(Club.id.in_(cids)).all()} if cids else {}
    out = []
    for p in pod:
        def nm(aid):
            a = amap.get(aid) if aid else None
            if not a:
                return None
            club = clubmap.get(a.club_id) if a.club_id else None
            return {"id": a.id, "name": a.full_name, "club": club.name if club else "—"}
        out.append({"category_id": p["category_id"], "category": p["category"],
                    "gold": nm(p["gold_id"]), "silver": nm(p["silver_id"]),
                    "bronze": [nm(b) for b in p["bronze_ids"] if nm(b)]})
    return out


@router.post("/api/tournaments/{tid}/documents/issue-podium")
def issue_podium(tid: int, db: Session = Depends(get_db), user: User = Depends(require_perm("documents.manage"))):
    """Wave 1: bulk diplomas for places 1-3 across finished categories."""
    from app.core.permissions import has_perm, user_roles
    t = db.get(Tournament, tid)
    if not t:
        raise HTTPException(404, "Not found")
    if not ("admin" in user_roles(db, user) or has_perm(db, user, "tournaments.manage_all")):
        if not (t.created_by == user.id and has_perm(db, user, "documents.manage")):
            raise HTTPException(403, "Foreign tournament")
    issued, skipped = [], 0
    for p in _podium(db, tid):
        spots = [(p["gold_id"], "1"), (p["silver_id"], "2")] + [(b, "3") for b in p["bronze_ids"]]
        for aid, place in spots:
            if not aid:
                continue
            d = _issue_one(db, aid, tid, "diploma", place, p["category"])
            if d is None:
                skipped += 1
            else:
                db.flush()
                issued.append({"athlete_id": aid, "place": place, "code": d.code})
    db.commit()
    return {"issued": issued, "skipped": skipped}


@router.get("/api/tournaments/{tid}/documents")
def list_documents(tid: int, db: Session = Depends(get_db), user: User = Depends(require_perm("documents.manage"))):
    """Wave 1: registry of issued documents for re-download (owner scope)."""
    t = db.get(Tournament, tid)
    if not t:
        raise HTTPException(404, "Not found")
    require_tournament_owner(tid, db, user)
    rows = db.query(Document).filter_by(tournament_id=tid).order_by(Document.id.desc()).all()
    amap = {a.id: a for a in db.query(Athlete).filter(
        Athlete.id.in_([d.athlete_id for d in rows if d.athlete_id])).all()} if rows else {}
    out = []
    for d in rows:
        try:
            payload = json.loads(d.payload or "{}")
        except ValueError:
            payload = {}
        a = amap.get(d.athlete_id) if d.athlete_id else None
        out.append({"code": d.code, "kind": d.kind, "athlete_id": d.athlete_id,
                    "athlete": a.full_name if a else "?", "place": payload.get("place", ""),
                    "category": payload.get("category", ""),
                    "template": payload.get("template", ""),
                    "at": str(d.created_at)})
    return out

@router.post("/api/documents/issue")
def issue_doc(athlete_id: int, tournament_id: int, kind: str = "participation",
              place: str = "", category: str = "",
              db: Session = Depends(get_db), user: User = Depends(require_perm("documents.manage"))):
    # P1: ownership must match the required permission. require_tournament_owner
    # checks tournaments.manage, which would force callers to hold BOTH perms;
    # a coach granted only documents.manage on their own tournament got 403.
    from app.core.permissions import has_perm, user_roles
    t = db.get(Tournament, tournament_id)
    if not t:
        raise HTTPException(404, "Not found")
    if not ("admin" in user_roles(db, user) or has_perm(db, user, "tournaments.manage_all")):
        if not (t.created_by == user.id and has_perm(db, user, "documents.manage")):
            raise HTTPException(403, "Foreign tournament")
    if kind not in ("participation", "diploma"):
        raise HTTPException(400, "Unknown document kind")
    reg = db.query(Registration).filter_by(tournament_id=tournament_id, athlete_id=athlete_id).first()
    if not reg:
        raise HTTPException(400, "Athlete is not registered in this tournament")
    # Wave 7: idempotent single issue (same helper as bulk podium) — a double
    # click or retry returns the existing code instead of minting duplicates.
    d = _issue_one(db, athlete_id, tournament_id, kind, place, category)
    if d is None:
        d = db.query(Document).filter_by(athlete_id=athlete_id,
                                        tournament_id=tournament_id, kind=kind).first()
    else:
        db.flush()
    db.commit()
    code = d.code
    # Wave 1: notify the athlete's coach (if someone else issued it).
    from app.services.notifications import emit_event, coach_of_athlete
    coach = coach_of_athlete(db, athlete_id)
    if coach and coach != user.id:
        emit_event(db, "document", [coach],
                   f"Документ готов: {kind} {place} (t{tournament_id})".strip(),
                   link=f"/tournaments/{tournament_id}?tab=results")
    return {"code": code}

@router.get("/api/documents/verify/{code}")
def verify_doc(code: str, db: Session = Depends(get_db)):
    d = db.query(Document).filter_by(code=code.upper()).first()
    if not d:
        return {"valid": False}
    a = db.get(Athlete, d.athlete_id) if d.athlete_id else None
    t = db.get(Tournament, d.tournament_id) if d.tournament_id else None
    try:
        payload = json.loads(d.payload or "{}")
    except ValueError:
        payload = {}
    return {"valid": True, "code": d.code, "kind": d.kind,
            "athlete": a.full_name if a else "?", "tournament": t.name if t else "?",
            "date": str(t.start_date) if t else "", **payload}

@router.get("/api/documents/{code}/certificate.pdf")
def certificate_pdf(code: str, db: Session = Depends(get_db)):
    import qrcode
    from reportlab.pdfgen import canvas
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import ImageReader
    d = db.query(Document).filter_by(code=code.upper()).first()
    if not d:
        raise HTTPException(404, "Document not found")
    v = verify_doc(code, db)
    from app.services.pdf_fonts import ensure_fonts
    from reportlab.pdfbase.pdfmetrics import stringWidth
    regular, bold = ensure_fonts()
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    c.setFont(bold, 26)
    c.drawCentredString(300, 700, "KWF")
    c.setFont(regular, 14)
    c.drawCentredString(300, 670, {"participation": "Certificate of Participation", "diploma": "Diploma"}.get(d.kind, d.kind))

    def _fit(text: str, font: str, size: int, max_w: float = 480, min_size: int = 10) -> int:
        # Wave 7: long names shrink instead of clipping past the page edge.
        while size > min_size and stringWidth(text, font, size) > max_w:
            size -= 1
        return size

    athlete, sub = v["athlete"], f"{v['tournament']}  ·  {v.get('category', '')}  ·  {v.get('place', '')}"
    c.setFont(bold, _fit(athlete, bold, 20))
    c.drawCentredString(300, 620, athlete)
    c.setFont(regular, _fit(sub, regular, 12))
    c.drawCentredString(300, 590, sub)
    qr = qrcode.make(f"verify:{d.code}")
    qbuf = io.BytesIO()
    qr.save(qbuf, format="PNG")
    qbuf.seek(0)
    c.drawImage(ImageReader(qbuf), 250, 400, 100, 100)
    c.setFont(regular, 10)
    c.drawCentredString(300, 380, f"Verify: {d.code}")
    c.showPage()
    c.save()
    return Response(buf.getvalue(), media_type="application/pdf",
                    headers={"Content-Disposition": f"attachment; filename=kwf-{d.code}.pdf"})

# ---------- audit log (§32) ----------

@router.get("/api/audit")
def list_audit(tournament_id: int | None = None, pg: dict = Depends(page_args), db: Session = Depends(get_db), user: User = Depends(require_perm("audit.view"))):
    """Audit log. Admin sees everything; organizers see only entries related
    to their own tournaments (or their own actions) — P0 fix for global leak
    where any organizer could read foreign tournaments' audit trail."""
    from app.core.permissions import has_perm, user_roles
    from sqlalchemy import or_, and_
    q = db.query(AuditLog).order_by(AuditLog.id.desc())
    if "admin" not in user_roles(db, user) and not has_perm(db, user, "tournaments.manage_all"):
        own_tids = {r[0] for r in db.query(Tournament.id).filter_by(created_by=user.id).all()}
        # Resolve match/registration audit rows back to their tournament so an
        # organizer sees foreign-actor events (e.g. referee's "finished fight")
        # on their own tournaments — but nothing from foreign tournaments.
        tids = [tournament_id] if tournament_id is not None else list(own_tids)
        if tournament_id is not None and tournament_id not in own_tids:
            raise HTTPException(403, "Foreign tournament")
        own_match_ids: set[int] = set()
        own_reg_ids: set[int] = set()
        if tids:
            bids = {r[0] for r in db.query(Bracket.id).filter(Bracket.tournament_id.in_(tids)).all()}
            if bids:
                own_match_ids = {r[0] for r in db.query(BracketMatch.id).filter(BracketMatch.bracket_id.in_(list(bids))).all()}
            own_reg_ids = {r[0] for r in db.query(Registration.id).filter(Registration.tournament_id.in_(tids)).all()}
        conds = [(AuditLog.actor_id == user.id)]
        if tids:
            conds.append(and_(AuditLog.entity == "tournament", AuditLog.entity_id.in_(tids)))
        if own_match_ids:
            conds.append(and_(AuditLog.entity == "match", AuditLog.entity_id.in_(list(own_match_ids))))
        if own_reg_ids:
            conds.append(and_(AuditLog.entity == "registration", AuditLog.entity_id.in_(list(own_reg_ids))))
        q = q.filter(or_(*conds))
    elif tournament_id is not None:
        q = q.filter(AuditLog.entity == "tournament", AuditLog.entity_id == tournament_id)
    rows, total = paginate(q, pg["limit"], pg["offset"])
    return envelope([{"id": r.id, "actor": r.actor_id, "action": r.action, "entity": r.entity,
             "entity_id": r.entity_id, "at": str(r.created_at)} for r in rows],
             total, pg["limit"], pg["offset"])

# ---------- tournament report (§46 Report/News Assistant) ----------

@router.get("/api/tournaments/{tid}/report")
def tournament_report(tid: int, lang: str = "ru", db: Session = Depends(get_db)):
    """Auto-generated post-tournament summary: stats + markdown draft for a news post.
    AI-free and factual — the organizer reviews before publishing."""
    kk = lang.lower().startswith("kk")
    t = db.get(Tournament, tid)
    if not t:
        raise HTTPException(404, "Not found")
    regs = db.query(Registration).filter_by(tournament_id=tid).all()
    athlete_ids = {r.athlete_id for r in regs}
    amap = {a.id: a for a in db.query(Athlete).filter(Athlete.id.in_(athlete_ids)).all()} if athlete_ids else {}
    athletes = list(amap.values())
    bids = [b.id for b in db.query(Bracket).filter_by(tournament_id=tid).all()]
    fights = db.query(BracketMatch).filter(BracketMatch.bracket_id.in_(bids)).all() if bids else []
    finished = [m for m in fights if m.status == "finished"]
    standings = _category_standings(db, tid)
    countries = sorted({a.country for a in athletes if a.country})
    cids = [a.club_id for a in athletes if a.club_id]
    clubmap = {c.id: c for c in db.query(Club).filter(Club.id.in_(cids)).all()} if cids else {}
    clubs = sorted({(clubmap[a.club_id].name if a.club_id and a.club_id in clubmap else "—") for a in athletes})
    club_gold: dict[str, int] = {}
    for s in standings:
        if s["champion_id"]:
            a = amap.get(s["champion_id"])
            club = clubmap[a.club_id].name if a and a.club_id and a.club_id in clubmap else "—"
            club_gold[club] = club_gold.get(club, 0) + 1
    lines = [f"# {t.name}: {'қорытынды' if kk else 'итоги'}", "",
             f"{'Қатысушылар' if kk else 'Участников'}: {len(athletes)} · {'Елдер' if kk else 'Стран'}: {len(countries)} · {'Клубтар' if kk else 'Клубов'}: {len(clubs)} · "
             f"{'Жекпе-жектер' if kk else 'Боёв'}: {len(finished)}/{len(fights)}", ""]
    for s in standings:
        lines.append(f"- {s['category']}: 🏆 {s['champion']}")
    stats = {"participants": len(athletes), "countries": len(countries), "clubs": len(clubs),
             "fights_total": len(fights), "fights_finished": len(finished),
             "club_gold": sorted(club_gold.items(), key=lambda x: -x[1])}
    return {"tournament": t.name, "stats": stats, "standings": standings, "markdown": "\n".join(lines)}

# ---------- Wave 2: club report (roster, tournaments, medals, points) ----------

def _club_report_data(db: Session, cid: int) -> dict:
    """Assemble club report from existing rows. No time series exists in the
    DB, so dynamics are honest totals (points/wins summed), not invented."""
    c = db.get(Club, cid)
    if not c:
        raise HTTPException(404, "Not found")
    athletes = db.query(Athlete).filter_by(club_id=cid).order_by(Athlete.points.desc()).all()
    aids = [a.id for a in athletes]
    regs = db.query(Registration).filter(Registration.athlete_id.in_(aids)).all() if aids else []
    tids = sorted({r.tournament_id for r in regs})
    tmap = {t.id: t for t in db.query(Tournament).filter(Tournament.id.in_(tids)).all()} if tids else {}
    gold = silver = bronze = 0
    per_tournament = []
    for tid in tids:
        t = tmap.get(tid)
        mine = {r.athlete_id for r in regs if r.tournament_id == tid}
        per_tournament.append({"id": tid, "name": t.name if t else "?",
                               "date": str(t.start_date) if t else "",
                               "participants": len(mine)})
        for p in _podium(db, tid):
            if p["gold_id"] in mine:
                gold += 1
            if p["silver_id"] in mine:
                silver += 1
            bronze += len([b for b in p["bronze_ids"] if b in mine])
    return {"club": {"id": c.id, "name": c.name, "country": c.country,
                     "city": c.city, "coach": c.coach_name},
            "roster": [{"id": a.id, "name": a.full_name, "points": a.points or 0,
                        "wins": a.wins or 0, "losses": a.losses or 0} for a in athletes],
            "tournaments": per_tournament,
            "stats": {"athletes": len(athletes),
                      "tournaments": len(tids),
                      "participations": len(regs),
                      "gold": gold, "silver": silver, "bronze": bronze,
                      "points": sum(a.points or 0 for a in athletes)}}


@router.get("/api/clubs/{cid}/report")
def club_report(cid: int, db: Session = Depends(get_db)):
    return _club_report_data(db, cid)


@router.get("/api/clubs/{cid}/report.pdf")
def club_report_pdf(cid: int, db: Session = Depends(get_db)):
    rep = _club_report_data(db, cid)
    rows = [[a["name"], a["points"], f"{a['wins']}-{a['losses']}"] for a in rep["roster"]]
    rows.append(["", "", ""])
    s = rep["stats"]
    rows += [[f"Tournaments: {s['tournaments']}", f"Entries: {s['participations']}", ""],
             [f"Gold: {s['gold']}", f"Silver: {s['silver']}", f"Bronze: {s['bronze']}"]]
    pdf = _sheet_pdf(f"Club report - {rep['club']['name']}", ["Athlete", "Points", "W-L"], rows or [["—", "", ""]])
    return Response(pdf, media_type="application/pdf",
                    headers={"Content-Disposition": f"attachment; filename=club{cid}-report.pdf"})


@router.get("/api/clubs/{cid}/report.xlsx")
def club_report_xlsx(cid: int, db: Session = Depends(get_db)):
    from openpyxl import Workbook
    rep = _club_report_data(db, cid)
    wb = Workbook()
    ws = wb.active
    ws.title = "Roster"
    ws.append(["Athlete", "Points", "Wins", "Losses"])
    for a in rep["roster"]:
        ws.append([a["name"], a["points"], a["wins"], a["losses"]])
    ws2 = wb.create_sheet("Tournaments")
    ws2.append(["Tournament", "Date", "Participants"])
    for t in rep["tournaments"]:
        ws2.append([t["name"], t["date"], t["participants"]])
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename=club{cid}-report.xlsx"})
